"""Small, deterministic KURAI presentation export without an external rendering service."""

from dataclasses import dataclass
from io import BytesIO
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


BACKGROUND = RGBColor(7, 12, 17)
PANEL = RGBColor(13, 23, 31)
TEXT = RGBColor(229, 237, 245)
MUTED = RGBColor(139, 160, 180)
ACCENT = RGBColor(0, 220, 234)


@dataclass(frozen=True)
class SlideSection:
    title: str
    bullets: tuple[str, ...]


def presentation_sections(text: str, *, max_sections: int = 12) -> list[SlideSection]:
    """Turn headings and paragraphs into bounded, readable presentation slides."""
    sections: list[SlideSection] = []
    current_title = "Ключевые выводы"
    current_bullets: list[str] = []

    def flush() -> None:
        nonlocal current_bullets
        if current_bullets and len(sections) < max_sections:
            sections.append(SlideSection(current_title[:100], tuple(current_bullets[:6])))
        current_bullets = []

    for raw_line in text.replace("\r\n", "\n").split("\n"):
        line = raw_line.strip()
        if not line:
            continue
        heading = re.match(r"^#{1,3}\s+(.+)$", line)
        if heading:
            flush()
            current_title = heading.group(1).strip() or "Ключевые выводы"
            continue
        bullet = re.sub(r"^(?:[-*•]|\d+[.)])\s+", "", line).strip()
        if bullet:
            current_bullets.append(bullet[:260])
            if len(current_bullets) == 6:
                flush()
    flush()
    return sections or [SlideSection("Ключевые выводы", (text.strip()[:500] or "Нет текста для экспорта.",))]


def _background(slide) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BACKGROUND
    accent = slide.shapes.add_shape(1, Inches(0.55), Inches(0.52), Inches(0.32), Inches(0.04))
    accent.fill.solid()
    accent.fill.fore_color.rgb = ACCENT
    accent.line.fill.background()


def _text_box(slide, *, left: float, top: float, width: float, height: float, text: str, size: int, color: RGBColor, bold: bool = False) -> None:
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.TOP
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = PP_ALIGN.LEFT
    paragraph.font.name = "Aptos"
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color


def _content_slide(prs: Presentation, section: SlideSection, number: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _background(slide)
    _text_box(slide, left=0.62, top=0.72, width=10.8, height=0.6, text=section.title, size=26, color=TEXT, bold=True)
    _text_box(slide, left=11.8, top=0.76, width=0.8, height=0.3, text=f"{number:02d}", size=10, color=ACCENT)
    panel = slide.shapes.add_shape(1, Inches(0.62), Inches(1.65), Inches(12.05), Inches(4.88))
    panel.fill.solid()
    panel.fill.fore_color.rgb = PANEL
    panel.line.fill.background()
    for index, bullet in enumerate(section.bullets):
        _text_box(slide, left=1.02, top=2.05 + index * 0.67, width=11.0, height=0.46, text=f"{index + 1:02d}  {bullet}", size=17, color=TEXT)


def build_presentation(*, title: str, text: str, sources=()) -> bytes:
    """Build a single private .pptx export. Text and links are never rendered as HTML."""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    cover = prs.slides.add_slide(prs.slide_layouts[6])
    _background(cover)
    _text_box(cover, left=0.62, top=1.35, width=10.8, height=0.45, text="KURAI / export", size=12, color=ACCENT)
    _text_box(cover, left=0.62, top=2.0, width=11.6, height=1.6, text=title[:120], size=38, color=TEXT, bold=True)
    _text_box(cover, left=0.62, top=4.45, width=10.8, height=0.5, text="Собрано из приватного ответа в рабочем пространстве", size=16, color=MUTED)

    for number, section in enumerate(presentation_sections(text), start=1):
        _content_slide(prs, section, number)

    source_rows = [(source.title[:180], source.url[:500]) for source in sources][:10]
    if source_rows:
        sources_slide = prs.slides.add_slide(prs.slide_layouts[6])
        _background(sources_slide)
        _text_box(sources_slide, left=0.62, top=0.72, width=10.8, height=0.6, text="Источники", size=26, color=TEXT, bold=True)
        for index, (source_title, url) in enumerate(source_rows):
            _text_box(sources_slide, left=1.02, top=1.65 + index * 0.48, width=11.1, height=0.33, text=f"{index + 1:02d}  {source_title} · {url}", size=11, color=MUTED)

    stream = BytesIO()
    prs.save(stream)
    return stream.getvalue()
