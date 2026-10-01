from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import Client
from pptx import Presentation

from workspace.models import Chat, Message


def test_owner_can_export_an_assistant_answer_to_a_private_presentation(db):
    owner = get_user_model().objects.create_user(email="slides@example.test", password="safe-test-password")
    client = Client()
    client.force_login(owner)
    chat = Chat.objects.create(owner=owner, title="Стратегия продукта")
    message = Message.objects.create(
        chat=chat,
        role=Message.Role.ASSISTANT,
        mode="slides",
        text="# Проблема\n\n- Слишком много ручной работы\n- Нет прозрачной цены\n\n# Решение\n\n- Единый каталог моделей\n- Проверяемые источники",
        parts=[{"type": "text", "text": "# Проблема\n\n- Слишком много ручной работы\n- Нет прозрачной цены\n\n# Решение\n\n- Единый каталог моделей\n- Проверяемые источники"}],
    )

    response = client.post(f"/api/v1/messages/{message.id}/exports/slides", data={"title": "Q4 plan"}, content_type="application/json")

    assert response.status_code == 200
    assert response["Content-Type"] == "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    assert response["Content-Disposition"].startswith("attachment;")
    deck = Presentation(BytesIO(response.content))
    assert len(deck.slides) == 3
    assert any("Решение" in shape.text for shape in deck.slides[2].shapes if hasattr(shape, "text"))


def test_slide_export_does_not_disclose_another_members_message(db):
    owner = get_user_model().objects.create_user(email="slides-owner@example.test", password="safe-test-password")
    stranger = get_user_model().objects.create_user(email="slides-stranger@example.test", password="safe-test-password")
    chat = Chat.objects.create(owner=owner, title="Private")
    message = Message.objects.create(chat=chat, role=Message.Role.ASSISTANT, text="Private answer", parts=[{"type": "text", "text": "Private answer"}])
    client = Client()
    client.force_login(stranger)

    response = client.post(f"/api/v1/messages/{message.id}/exports/slides", data={}, content_type="application/json")

    assert response.status_code == 404
